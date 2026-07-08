"""Convert extracted invoice data into a formatted Excel workbook.

Builds a workbook with two sheets:
- "Invoices": one row per PDF with the header-level fields.
- "Line Items": one row per line item, linked back to its invoice.
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# Header-level columns shown on the "Invoices" sheet.
SUMMARY_COLUMNS = [
    ("file", "File"),
    ("vendor_name", "Vendor"),
    ("vendor_address", "Vendor Address"),
    ("buyer_name", "Buyer"),
    ("buyer_address", "Buyer Address"),
    ("invoice_number", "Invoice #"),
    ("invoice_date", "Invoice Date"),
    ("due_date", "Due Date"),
    ("subtotal", "Subtotal"),
    ("tax_amount", "Tax"),
    ("total_amount", "Total"),
    ("currency", "Currency"),
    ("payment_terms", "Payment Terms"),
    ("notes", "Notes"),
    ("extraction_method", "Method"),
]

LINE_ITEM_COLUMNS = [
    ("_invoice_ref", "Invoice"),
    ("description", "Description"),
    ("quantity", "Quantity"),
    ("unit_price", "Unit Price"),
    ("amount", "Amount"),
]

_HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_THIN = Side(style="thin", color="D9D9D9")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def _autofit(ws, headers, max_width=60):
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        longest = len(str(header))
        for cell in ws[letter]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_width, max(10, longest + 2))


def _display_ref(record, index):
    """Human-friendly label for an invoice (used to link line items)."""
    return (
        record.get("invoice_number")
        or record.get("original_filename")
        or record.get("file")
        or f"Invoice {index + 1}"
    )


def build_workbook(records):
    """Build an Excel workbook from a list of extracted-invoice dicts.

    Args:
        records: list[dict] as returned by extract_invoice_data().

    Returns:
        openpyxl.Workbook
    """
    wb = Workbook()

    # ── Invoices sheet ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Invoices"
    summary_headers = [label for _, label in SUMMARY_COLUMNS]
    ws.append(summary_headers)

    for record in records:
        row = []
        for key, _ in SUMMARY_COLUMNS:
            value = record.get(key)
            if key == "file":
                value = record.get("original_filename") or record.get("file")
            row.append(value if value is not None else "")
        ws.append(row)

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(summary_headers) + 1):
            ws.cell(row=r, column=c).border = _BORDER
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    _style_header(ws, len(summary_headers))
    _autofit(ws, summary_headers)

    # ── Line Items sheet ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Line Items")
    line_headers = [label for _, label in LINE_ITEM_COLUMNS]
    ws2.append(line_headers)

    any_items = False
    for index, record in enumerate(records):
        items = record.get("line_items") or []
        ref = _display_ref(record, index)
        for item in items:
            if not isinstance(item, dict):
                continue
            any_items = True
            ws2.append([
                ref,
                item.get("description", ""),
                item.get("quantity", "") if item.get("quantity") is not None else "",
                item.get("unit_price", "") if item.get("unit_price") is not None else "",
                item.get("amount", "") if item.get("amount") is not None else "",
            ])

    if not any_items:
        ws2.append(["(no line items extracted)", "", "", "", ""])

    for r in range(2, ws2.max_row + 1):
        for c in range(1, len(line_headers) + 1):
            ws2.cell(row=r, column=c).border = _BORDER
            ws2.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    _style_header(ws2, len(line_headers))
    _autofit(ws2, line_headers)

    return wb


def workbook_to_bytes(records):
    """Return the .xlsx file contents as bytes."""
    wb = build_workbook(records)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
